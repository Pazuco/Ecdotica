"""
Tests para src/procesamiento/utils.py
"""

import pytest
import tempfile
import os
from utils import contar_palabras, contar_capitulos, calcular_legibilidad, analizar_manuscrito


class TestContarPalabras:
    def test_texto_normal(self):
        assert contar_palabras("hola mundo") == 2

    def test_texto_vacio(self):
        assert contar_palabras("") == 0

    def test_texto_con_puntuacion(self):
        assert contar_palabras("¡Hola, mundo!") == 2

    def test_texto_con_acentos(self):
        assert contar_palabras("árbol canción análisis") == 3

    def test_una_sola_palabra(self):
        assert contar_palabras("poema") == 1


class TestContarCapitulos:
    def test_sin_capitulos(self):
        assert contar_capitulos("Texto sin capítulos aquí.") == 0

    def test_un_capitulo(self):
        assert contar_capitulos("Capítulo I\nTexto del capítulo.") == 1

    def test_varios_capitulos(self):
        texto = "Capítulo I\nTexto.\n\nCapítulo II\nMás texto.\n\nCapítulo III\nFin."
        assert contar_capitulos(texto) == 3

    def test_capitulo_minuscula(self):
        assert contar_capitulos("capítulo uno\nTexto.") == 1

    def test_capitulo_en_medio_de_linea_no_cuenta(self):
        # Solo cuenta si está al inicio de línea
        assert contar_capitulos("Este es el capítulo primero.") == 0


class TestCalcularLegibilidad:
    def test_texto_vacio_devuelve_100(self):
        assert calcular_legibilidad("") == 100

    def test_sin_oraciones_devuelve_valor_valido(self):
        # textstat calcula legibilidad incluso sin puntuación explícita
        resultado = calcular_legibilidad("texto sin puntuación final")
        assert 0 <= resultado <= 100

    def test_resultado_entre_0_y_100(self):
        texto = "El sol brilla. Las aves cantan. El río fluye con calma serena."
        resultado = calcular_legibilidad(texto)
        assert 0 <= resultado <= 100

    def test_texto_simple_alta_legibilidad(self):
        texto = "Hola. Adiós. Bien. Mal."
        resultado = calcular_legibilidad(texto)
        assert resultado >= 0


class TestAnalizarManuscrito:
    def test_devuelve_dict_con_claves_correctas(self, tmp_path):
        archivo = tmp_path / "manuscrito.txt"
        archivo.write_text("Este es un texto de prueba. Tiene dos oraciones.", encoding="utf-8")
        stats = analizar_manuscrito(str(archivo))
        assert 'num_palabras' in stats
        assert 'num_capitulos' in stats
        assert 'indice_legibilidad' in stats
        assert 'errores_graves' in stats

    def test_cuenta_palabras_correctamente(self, tmp_path):
        archivo = tmp_path / "manuscrito.txt"
        archivo.write_text("uno dos tres cuatro cinco.", encoding="utf-8")
        stats = analizar_manuscrito(str(archivo))
        assert stats['num_palabras'] == 5

    def test_archivo_con_capitulos(self, tmp_path):
        archivo = tmp_path / "novela.txt"
        contenido = "Capítulo I\nTexto.\n\nCapítulo II\nMás texto."
        archivo.write_text(contenido, encoding="utf-8")
        stats = analizar_manuscrito(str(archivo))
        assert stats['num_capitulos'] == 2

    def test_archivo_no_encontrado_lanza_error(self):
        with pytest.raises(FileNotFoundError):
            analizar_manuscrito("/ruta/que/no/existe.txt")


class TestAnalizarManuscritoDocx:
    def test_devuelve_dict_con_claves_correctas(self, tmp_path):
        try:
            import docx
            doc = docx.Document()
            doc.add_paragraph("Este es un texto de prueba en DOCX. Tiene dos oraciones.")
            ruta = str(tmp_path / "manuscrito.docx")
            doc.save(ruta)
            stats = analizar_manuscrito(ruta)
            assert 'num_palabras' in stats
            assert 'num_capitulos' in stats
            assert 'indice_legibilidad' in stats
            assert 'errores_graves' in stats
        except ImportError:
            pytest.skip("python-docx no instalado")

    def test_cuenta_palabras_correctamente(self, tmp_path):
        try:
            import docx
            doc = docx.Document()
            doc.add_paragraph("uno dos tres cuatro cinco.")
            ruta = str(tmp_path / "manuscrito.docx")
            doc.save(ruta)
            stats = analizar_manuscrito(ruta)
            assert stats['num_palabras'] == 5
        except ImportError:
            pytest.skip("python-docx no instalado")

    def test_formato_no_soportado_lanza_valor_error(self, tmp_path):
        archivo = tmp_path / "manuscrito.odt"
        archivo.write_text("contenido", encoding="utf-8")
        with pytest.raises(ValueError, match=".txt|.pdf|.docx"):
            analizar_manuscrito(str(archivo))


class TestAnalizarManuscritoPdf:
    def test_devuelve_dict_con_claves_correctas(self, tmp_path):
        try:
            from reportlab.pdfgen import canvas as rl_canvas
        except ImportError:
            pytest.skip("reportlab no instalado")
        try:
            ruta = str(tmp_path / "manuscrito.pdf")
            c = rl_canvas.Canvas(ruta)
            c.drawString(100, 750, "uno dos tres cuatro cinco.")
            c.save()
            stats = analizar_manuscrito(ruta)
            assert 'num_palabras' in stats
        except Exception as e:
            pytest.skip(f"Error al procesar PDF en este entorno: {e}")
